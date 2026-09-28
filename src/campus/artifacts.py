import hashlib
import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from campus.models import CampusError, now
from campus.providers.local import git
from campus.security import has_secret, inside, redact

EXCLUDED = {
    ".git",
    ".svn",
    "node_modules",
    ".venv",
    "venv",
    "build",
    "dist",
    "target",
    "bin",
    "obj",
    "__pycache__",
    ".campus",
    ".codex",
    ".agents",
    ".ssh",
    ".aws",
    ".azure",
    ".config",
    "auth",
    "credentials",
    "cookies",
    "debug",
    "reports",
    "artifacts",
    ".pytest_cache",
    ".ruff_cache",
    "coverage",
}
TEXT_EXTENSIONS = {
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".kt",
    ".kts",
    ".c",
    ".cpp",
    ".h",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".dart",
    ".swift",
    ".html",
    ".css",
    ".scss",
    ".sql",
    ".md",
    ".txt",
    ".json",
    ".toml",
    ".yaml",
    ".yml",
    ".xml",
    ".csv",
    ".svg",
    ".sln",
    ".csproj",
    ".gradle",
    ".properties",
    ".sh",
    ".ps1",
    ".bat",
    ".r",
    ".ipynb",
}
DENY_NAME = re.compile(
    r"(^\.env(?:\.|$)|credential|secret|password|passwd|cookie|storage.?state|(^|[_.-])(?:token|auth|private.?key)(?:[_.-]|$)|id_rsa|id_ed25519|\.npmrc$|\.pypirc$|\.netrc$)",
    re.I,
)


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def archive_entry(archive, name: str, data: bytes):
    info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o100644 << 16
    archive.writestr(info, data)


def package_repository(repository: Path, output: Path) -> dict:
    repository = repository.resolve()
    if not repository.is_dir() or not (repository / ".git").exists():
        raise CampusError("ZIP source must be a local Git repository")
    # Tracked source only. Ignored and untracked personal files never enter packages.
    names = git(repository, "ls-files", "-z").split("\0")
    selected, excluded, total = [], [], 0
    for name in sorted(n for n in names if n):
        rel = Path(name)
        path = repository / rel
        if (
            rel.is_absolute()
            or ".." in rel.parts
            or any(p.lower() in EXCLUDED or DENY_NAME.search(p) for p in rel.parts)
        ):
            excluded.append({"path": redact(name), "reason": "sensitive/cache path"})
            continue
        if path.is_symlink() or any(
            p.is_symlink() or p.is_junction() for p in path.parents if p.is_relative_to(repository)
        ):
            excluded.append({"path": redact(name), "reason": "symbolic link or junction"})
            continue
        if (
            not path.is_file()
            or not path.resolve().is_relative_to(repository)
            or path.resolve() == output.resolve()
        ):
            excluded.append({"path": redact(name), "reason": "not an eligible source file"})
            continue
        if path.suffix.lower() not in TEXT_EXTENSIONS and path.name not in {
            "Dockerfile",
            "Makefile",
            "LICENSE",
            ".gitignore",
            ".editorconfig",
        }:
            excluded.append(
                {"path": redact(name), "reason": "binary/unknown format requires manual review"}
            )
            continue
        if path.stat().st_size > 5_000_000:
            excluded.append({"path": redact(name), "reason": "file exceeds 5 MB"})
            continue
        data = path.read_bytes()
        if has_secret(data) or b"\0" in data:
            excluded.append({"path": redact(name), "reason": "possible secret or nontext content"})
            continue
        total += len(data)
        if total > 100_000_000:
            raise CampusError("Source package exceeds 100 MB safety limit")
        selected.append((rel.as_posix(), data))
    if not selected:
        raise CampusError("No eligible tracked source files to package")
    if output.exists():
        raise CampusError("Artifact already exists; choose a new name")
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output, "x") as archive:
        for name, data in selected:
            archive_entry(archive, name, data)
    return {
        "path": str(output),
        "sha256": file_hash(output),
        "included": [n for n, _ in selected],
        "file_hashes": {n: hashlib.sha256(data).hexdigest() for n, data in selected},
        "excluded": excluded,
        "validation": "CRC checked" if validate_artifact(output) else "FAILED",
        "warning": "Source-only package. Exclusions must be reviewed against assignment requirements; secret detection is defense in depth, not proof that arbitrary source text contains no secret.",
    }


def generate_document(text: str, output: Path) -> dict:
    if output.exists():
        raise CampusError("Artifact already exists; choose a new name")
    if has_secret(text.encode()):
        raise CampusError("Document generation blocked: input may contain a secret")
    text = redact(text)
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = output.suffix.lower()
    if suffix in {".md", ".txt"}:
        output.write_text(text, encoding="utf-8")
    elif suffix == ".pdf":
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        document = SimpleDocTemplate(
            str(output), invariant=1, title="CAMPUS local document", author="CAMPUS"
        )
        style = getSampleStyleSheet()["BodyText"]
        story = []
        for line in text.splitlines():
            story.append(Paragraph(escape(line) or "&#160;", style))
            story.append(Spacer(1, 4))
        document.build(story)
    elif suffix == ".docx":
        from docx import Document

        document = Document()
        document.core_properties.author = "CAMPUS"
        document.core_properties.created = datetime(2000, 1, 1)
        document.core_properties.modified = datetime(2000, 1, 1)
        for line in text.splitlines():
            document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(output, "x") as target:
            for entry in sorted(source.namelist()):
                archive_entry(target, entry, source.read(entry))
    else:
        raise CampusError("Supported documents: .md, .txt, .pdf, .docx")
    return {
        "path": str(output),
        "sha256": file_hash(output),
        "validation": "PASSED" if validate_artifact(output) else "FAILED",
    }


def validate_artifact(path: Path) -> bool:
    try:
        if path.suffix in {".zip", ".docx"}:
            with zipfile.ZipFile(path) as archive:
                if archive.testzip() is not None:
                    return False
                if any(Path(n).is_absolute() or ".." in Path(n).parts for n in archive.namelist()):
                    return False
                if path.suffix == ".docx" and "word/document.xml" not in archive.namelist():
                    return False
        elif path.suffix == ".pdf":
            from pypdf import PdfReader

            return len(PdfReader(path).pages) > 0
        else:
            path.read_text(encoding="utf-8")
        return path.stat().st_size > 0
    except Exception:
        return False


def artifact_manifest(
    directory: Path, artifacts: list[dict], filename="manifest.json", **context
) -> Path:
    from campus.security import clean

    target = inside(directory, directory / filename)
    if target.exists():
        raise CampusError("Manifest already exists")
    target.write_text(
        json.dumps(
            clean({"created_at": now(), **context, "artifacts": artifacts}),
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return target


def read_document(path: Path, limit=100000) -> str:
    if not path.is_file() or path.stat().st_size > 20_000_000:
        raise CampusError("Document missing or larger than 20 MB")
    if path.suffix.lower() in {".md", ".txt", ".csv"}:
        return redact(path.read_text(encoding="utf-8", errors="replace")[:limit])
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise CampusError("Encrypted PDF requires manual access")
        chunks = []
        for page in reader.pages[:100]:
            chunks.append((page.extract_text() or "")[:limit])
            if sum(map(len, chunks)) >= limit:
                break
        return redact("\n".join(chunks)[:limit])
    if path.suffix.lower() == ".docx":
        from docx import Document

        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50_000_000:
                raise CampusError("DOCX uncompressed content exceeds safety limit")
        return redact("\n".join(p.text for p in Document(path).paragraphs)[:limit])
    raise CampusError("Unsupported document format; no macros or embedded content will be executed")
